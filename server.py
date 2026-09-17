from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlparse
import csv
import io
import json
import openpyxl

ROOT = Path(__file__).parent

class Handler(SimpleHTTPRequestHandler):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=str(ROOT), **kwargs)

    def do_POST(self):
        if urlparse(self.path).path != '/api/upload':
            self.send_error(404)
            return
        length = int(self.headers.get('Content-Length', '0'))
        payload = self.rfile.read(length)
        name = self.headers.get('X-File-Name', 'upload.xlsx').lower()
        try:
            if name.endswith(('.xlsx', '.xlsm', '.xltx', '.xltm')):
                workbook = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
                sheet = workbook[workbook.sheetnames[0]]
                rows = list(sheet.iter_rows(values_only=True))
                workbook.close()
                output = io.StringIO()
                csv.writer(output, lineterminator='\n').writerows(rows)
                body = output.getvalue().encode('utf-8-sig')
            else:
                body = payload
            self.send_response(200)
            self.send_header('Content-Type', 'text/csv; charset=utf-8')
            self.send_header('Access-Control-Allow-Origin', '*')
            self.send_header('Content-Length', str(len(body)))
            self.end_headers()
            self.wfile.write(body)
        except Exception as exc:
            body = json.dumps({'error': str(exc)}, ensure_ascii=False).encode('utf-8')
            self.send_response(400)
            self.send_header('Content-Type', 'application/json; charset=utf-8')
            self.send_header('Access-Control-Allow-Origin', '*')
            self.send_header('Content-Length', str(len(body)))
            self.end_headers()
            self.wfile.write(body)

if __name__ == '__main__':
    print('TSC server running at http://localhost:8765')
    ThreadingHTTPServer(('localhost', 8765), Handler).serve_forever()
